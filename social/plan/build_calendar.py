# Hotcoin marketing calendar, Oct 19 - Nov 15 2026. EN + 中文 tabs.
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

OUT = '/home/user/sandbox/social/plan/hotcoin-oct-nov-2026-calendar.xlsx'
INK, LIME, PANEL, LINE = '0B0E11', 'B8F26A', 'F4F5F2', 'D9DDD5'
ORANGE, GOLD, BLUE, GREY = 'FFE3C2', 'E8F7D2', 'E3ECFF', 'EEF0F2'
thin = Side(style='thin', color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

T = {
 'en': dict(
  sheet='Calendar (EN)',
  title='Hotcoin · Marketing calendar · Oct 19 – Nov 15, 2026',
  sub='Halloween theme Oct 19–31 · Ronaldinho global affiliate teasers Nov 1–4, reveal Nov 5 · all vouchers = futures trial vouchers (USD)',
  gridH='1 · Weekly plan by workstream', ws='Workstream',
  weeks=['Wk 1 · Oct 19–25', 'Wk 2 · Oct 26 – Nov 1', 'Wk 3 · Nov 2–8', 'Wk 4 · Nov 9–15'],
  rows=[
   ('Theme', ORANGE, ['🎃 Spooky season starts (light: Halloween look on regular posts)', '🎃 Halloween peak · Oct 31 = Halloween + Bitcoin whitepaper 18 yrs', '⚽ Ronaldinho reveal week', '⚽ Ronaldinho follow-up · 11.11']),
   ('Ronaldinho (global affiliate)', GOLD, ['No official promo in October (per agency)', 'Nov 1 (Sun): teaser 1 of 5', 'Nov 2–4: teasers 2–4 · Nov 5 (Thu): REVEAL — landing page post + media release + synced superstar social post', 'Sustain: official Brazil-shoot creatives only']),
   ('X giveaways (3 / week)', BLUE, ['2 themed: Halloween "trick or trade" · earnings season\n1 new users only', '2 themed: Fed day (Oct 28) · Halloween / whitepaper\n1 new users only', '2 themed: Ronaldinho reveal (if in scope) · Diwali\n1 new users only', '2 themed: 11.11 · market theme of the week\n1 new users only']),
   ('Paid ads', GREY, ['Spooky season campaign · Meta (main) · ChatGPT ads (test) · TG ads (explore)', 'Spooky season campaign peak · ends Oct 31', 'Switch to Ronaldinho creatives from Nov 5 (after review)', 'Ronaldinho campaign']),
   ('Discord', GREY, ['6 daily events (tweet raids, message boost, platform interaction, 1 trade/week…) in Halloween style · 2 weekly games · live trading lesson', 'Same · Halloween finale event', 'Same · Ronaldinho reveal watch-party / quiz', 'Same']),
  ],
  keyH='2 · Key dates', kcols=['Date', 'Day', 'Event', 'What we do'],
  keys=[
   ('Oct 19', 'Mon', 'Halloween theme starts', 'Light Halloween look on all channels; launch spooky season ads'),
   ('~Oct 21', 'Wed', 'Tesla earnings (est.) — big-tech earnings season', 'TradFi / stock futures post'),
   ('Oct 27–28', 'Tue–Wed', 'FOMC · decision Oct 28, 2:00 PM ET', 'Keep Oct 28 market-first, not costume-first'),
   ('~Oct 29', 'Thu', 'Apple, Amazon earnings (est.)', 'TradFi post'),
   ('Oct 31', 'Sat', 'Halloween + Bitcoin whitepaper 18 yrs + end of Uptober', 'Peak Halloween day · no Ronaldinho reveal here'),
   ('Nov 1', 'Sun', 'Ronaldinho teaser 1 · Moonvember starts', 'Teaser video 1 of 5'),
   ('Nov 2–4', 'Mon–Wed', 'Ronaldinho teasers 2–4', 'One teaser per day; Nov 3 = US midterms, no political takes'),
   ('Nov 5', 'Thu', 'Ronaldinho REVEAL', 'Landing page post + media release + synced superstar social post'),
   ('Nov 6', 'Fri', 'US October jobs report (usually 1st Friday)', 'Macro post'),
   ('Nov 8', 'Sun', 'Diwali (festival Nov 6–10)', 'Greeting post'),
   ('Nov 11', 'Wed', 'Singles\' Day 11.11', 'Promo post / giveaway theme'),
  ],
  budH='3 · Weekly voucher budget (USD, futures trial)', bcols=['Item', 'Wk 1', 'Wk 2', 'Wk 3', 'Wk 4', 'Total'],
  bud=[
   ('X themed giveaways (2 × $100)', 200, 'From brief: 2 weekly themed giveaways, $100 each to 10 random winners'),
   ('X new-users giveaway (1 × $100)', 100, 'From brief: 1 weekly giveaway for new users only, $100'),
   ('Discord 6-day programme (top 5 × $10)', 50, 'From Discord announcement: top 5 × $10'),
   ('Discord games (DLS + eFootball, $10 each)', 20, 'Assumes one winner per game per week'),
   ('Trading lesson with Mellst (max)', 70, 'Range $45–70 per week; pending approval from Loris & Elena'),
  ],
  tot='Total', note='Blue numbers are inputs — edit them and totals update. Paid ads budget not included (TBD). Ronaldinho creatives: official Brazil shoot only, no AI spokesperson images; review via Siri.',
 ),
 'cn': dict(
  sheet='日程 (中文)',
  title='Hotcoin · 营销日程 · 2026 年 10 月 19 日 – 11 月 15 日',
  sub='万圣节主题 10/19–10/31 · 小罗全球代言预热 11/1–11/4，11/5 正式官宣 · 所有奖励均为合约体验金（美元）',
  gridH='1 · 各板块周计划', ws='板块',
  weeks=['第 1 周 · 10/19–10/25', '第 2 周 · 10/26–11/1', '第 3 周 · 11/2–11/8', '第 4 周 · 11/9–11/15'],
  rows=[
   ('主题', ORANGE, ['🎃 万圣节季开始（轻度：常规内容加入万圣节元素）', '🎃 万圣节高峰 · 10/31 = 万圣节 + 比特币白皮书 18 周年', '⚽ 小罗官宣周', '⚽ 小罗持续推广 · 双十一']),
   ('小罗（全球代言）', GOLD, ['10 月无官方大规模推广（代理商要求）', '11/1（周日）：预热视频 1/5', '11/2–11/4：预热 2–4 · 11/5（周四）：正式官宣 —— 落地页帖子 + 媒体新闻稿 + 球星社媒同步发布', '持续推广：仅使用巴西官方拍摄素材']),
   ('X 抽奖（每周 3 场）', BLUE, ['2 场主题：万圣节"不交易就捣蛋" · 财报季\n1 场仅限新用户', '2 场主题：美联储议息日（10/28）· 万圣节 / 白皮书\n1 场仅限新用户', '2 场主题：小罗官宣（如授权允许）· 排灯节\n1 场仅限新用户', '2 场主题：双十一 · 本周市场热点\n1 场仅限新用户']),
   ('付费广告', GREY, ['万圣节季广告 · Meta（主力）· ChatGPT 广告（测试）· TG 广告（探索）', '万圣节广告高峰 · 10/31 结束', '11/5 起切换为小罗素材（审核后）', '小罗广告投放']),
   ('Discord', GREY, ['6 项每日活动（推特冲榜、消息互动、平台互动、每周一笔交易等）万圣节风格 · 每周 2 场游戏 · 直播交易课堂', '同上 · 万圣节压轴活动', '同上 · 小罗官宣观看派对 / 竞猜', '同上']),
  ],
  keyH='2 · 关键日期', kcols=['日期', '星期', '事件', '我们的动作'],
  keys=[
   ('10/19', '周一', '万圣节主题开始', '全渠道加入轻度万圣节元素；上线万圣节广告'),
   ('约 10/21', '周三', '特斯拉财报（预计）· 科技股财报季', 'TradFi / 美股合约内容'),
   ('10/27–10/28', '周二–周三', '美联储议息会议 · 10/28 美东 14:00 公布', '10/28 以市场内容为主，弱化万圣节'),
   ('约 10/29', '周四', '苹果、亚马逊财报（预计）', 'TradFi 内容'),
   ('10/31', '周六', '万圣节 + 比特币白皮书 18 周年 + Uptober 收官', '万圣节高峰日 · 不在当天官宣小罗'),
   ('11/1', '周日', '小罗预热 1 · Moonvember 开始', '预热视频 1/5'),
   ('11/2–11/4', '周一–周三', '小罗预热 2–4', '每天一条预热；11/3 美国中期选举，不发表政治观点'),
   ('11/5', '周四', '小罗正式官宣', '落地页帖子 + 媒体新闻稿 + 球星社媒同步发布'),
   ('11/6', '周五', '美国 10 月非农就业数据（通常为首个周五）', '宏观内容'),
   ('11/8', '周日', '排灯节（节期 11/6–11/10）', '节日祝福帖'),
   ('11/11', '周三', '双十一', '促销帖 / 抽奖主题'),
  ],
  budH='3 · 每周体验金预算（美元，合约体验金）', bcols=['项目', '第 1 周', '第 2 周', '第 3 周', '第 4 周', '合计'],
  bud=[
   ('X 主题抽奖（2 × $100）', 200, '来源：每周 2 场主题抽奖，每场 $100 分给 10 位随机用户'),
   ('X 新用户抽奖（1 × $100）', 100, '来源：每周 1 场仅限新用户抽奖，$100'),
   ('Discord 六日活动（前 5 名 × $10）', 50, '来源：Discord 公告，前 5 名各 $10'),
   ('Discord 游戏（DLS + eFootball，各 $10）', 20, '假设每个游戏每周 1 位冠军'),
   ('Mellst 交易课堂（上限）', 70, '每周 $45–70；待 Loris 与 Elena 审批'),
  ],
  tot='合计', note='蓝色数字为输入项，修改后合计自动更新。未含付费广告预算（待定）。小罗素材：仅用巴西官方拍摄素材，不使用 AI 生成的代言人形象；需经 Siri 审核。',
 ),
}

wb = Workbook(); wb.remove(wb.active)
for k in ('en', 'cn'):
    L = T[k]; ws = wb.create_sheet(L['sheet'])
    F = 'Arial'
    ws.sheet_view.showGridLines = False
    for col, w in zip('ABCDEF', [30, 36, 36, 36, 36, 14]):
        ws.column_dimensions[col].width = w
    ws['A1'] = L['title']; ws['A1'].font = Font(name=F, size=16, bold=True, color=INK)
    ws['A2'] = L['sub']; ws['A2'].font = Font(name=F, size=10, color='6B7280')
    r = 4
    ws.cell(r, 1, L['gridH']).font = Font(name=F, size=12, bold=True, color=INK); r += 1
    hdr = [L['ws']] + L['weeks']
    for c, v in enumerate(hdr, 1):
        cell = ws.cell(r, c, v); cell.font = Font(name=F, bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor=INK)
        cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True); cell.border = BOX
    ws.row_dimensions[r].height = 22; grid_hdr = r; r += 1
    for name, fill, cells in L['rows']:
        a = ws.cell(r, 1, name); a.font = Font(name=F, bold=True, color=INK); a.fill = PatternFill('solid', fgColor=fill)
        a.alignment = Alignment(vertical='top', wrap_text=True); a.border = BOX
        for c, v in enumerate(cells, 2):
            cell = ws.cell(r, c, v); cell.font = Font(name=F, size=10, color=INK)
            cell.fill = PatternFill('solid', fgColor=fill); cell.alignment = Alignment(vertical='top', wrap_text=True); cell.border = BOX
        ws.row_dimensions[r].height = 64; r += 1
    r += 1
    ws.cell(r, 1, L['keyH']).font = Font(name=F, size=12, bold=True, color=INK); r += 1
    for c, v in enumerate(L['kcols'], 1):
        cell = ws.cell(r, c, v); cell.font = Font(name=F, bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor=INK); cell.border = BOX
    r += 1
    for d in L['keys']:
        hot = 'Ronaldinho' in d[2] or '小罗' in d[2]
        for c, v in enumerate(d, 1):
            cell = ws.cell(r, c, v); cell.font = Font(name=F, size=10, color=INK, bold=(c == 1 or hot))
            cell.alignment = Alignment(vertical='top', wrap_text=True); cell.border = BOX
            if hot: cell.fill = PatternFill('solid', fgColor=GOLD)
            elif d[0].startswith(('Oct 31', '10/31', 'Oct 19', '10/19')): cell.fill = PatternFill('solid', fgColor=ORANGE)
        r += 1
    r += 1
    ws.cell(r, 1, L['budH']).font = Font(name=F, size=12, bold=True, color=INK); r += 1
    for c, v in enumerate(L['bcols'], 1):
        cell = ws.cell(r, c, v); cell.font = Font(name=F, bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor=INK); cell.border = BOX
        if c > 1: cell.alignment = Alignment(horizontal='right')
    r += 1; first = r
    for name, amt, src in L['bud']:
        a = ws.cell(r, 1, name); a.font = Font(name=F, size=10); a.border = BOX; a.comment = Comment(src, 'Hotcoin marketing')
        for c in range(2, 6):
            cell = ws.cell(r, c, amt if c == 2 else f'={ws.cell(r, 2).coordinate}')
            cell.font = Font(name=F, size=10, color='0000FF' if c == 2 else '000000')
            cell.number_format = '$#,##0'; cell.border = BOX
        t = ws.cell(r, 6, f'=SUM(B{r}:E{r})'); t.font = Font(name=F, size=10, bold=True); t.number_format = '$#,##0'; t.border = BOX
        r += 1
    last = r - 1
    a = ws.cell(r, 1, L['tot']); a.font = Font(name=F, bold=True, color=INK); a.fill = PatternFill('solid', fgColor=LIME); a.border = BOX
    for c in range(2, 7):
        col = ws.cell(r, c).column_letter
        cell = ws.cell(r, c, f'=SUM({col}{first}:{col}{last})'); cell.font = Font(name=F, bold=True, color=INK)
        cell.fill = PatternFill('solid', fgColor=LIME); cell.number_format = '$#,##0'; cell.border = BOX
    r += 2
    ws.cell(r, 1, L['note']).font = Font(name=F, size=9, italic=True, color='6B7280')
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[r].height = 30
    ws.freeze_panes = ws.cell(grid_hdr + 1, 2)
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT); print(OUT)
